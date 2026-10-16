#!/usr/bin/env python3
"""Fill missing postal codes from OneMap, only where the match is unambiguous.

  python _directory-sync/enrich_postal.py            # look up (network) and update listings
  python _directory-sync/enrich_postal.py --no-lookup  # use cached results only

1. Collects every published community listing and every review-queue record
   that has a street address (block + road) but no postal code.
2. Looks each address up on OneMap and caches the outcome in
   state/onemap-postal-cache.csv. A code is accepted only when OneMap returns
   the same block and road within 200 m of the listing's own map position and
   all matches agree on one code. Everything else is left blank.
3. Updates published community listings with the postal code and the area,
   sector and district it implies, and records the OneMap attribution.
   (The monthly sync applies the same cache to review-queue records.)
"""
import argparse
import datetime as dt
import sys
from pathlib import Path
from zoneinfo import ZoneInfo

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import openpyxl  # noqa: E402

from dsync import SGT_TZ  # noqa: E402
from dsync.geocode import ONEMAP_ATTRIBUTION, address_key, postal_for_address, throttle  # noqa: E402
from dsync.normalise import metres_between  # noqa: E402
from dsync.workbook import SHEET, _safe_commit, postal_lookup, read_directory, update_records  # noqa: E402
from sync import STATE, WORKBOOK, community_path, read_csv, write_csv  # noqa: E402

CACHE = STATE / "onemap-postal-cache.csv"
CACHE_FIELDS = ["address_key", "latitude", "longitude", "postal", "outcome", "checked_on"]
RETRY_DAYS = 90


def load_cache():
    return read_csv(CACHE)


def cached_postal(cache, address, lat, lon):
    key = address_key(address)
    if not key:
        return ""
    for c in cache:
        if c["address_key"] == key and c["outcome"] == "matched":
            d = metres_between(lat, lon, c["latitude"], c["longitude"])
            if d is not None and d <= 200:
                return c["postal"]
    return ""


def lookup(targets, cache, today):
    fresh = {}
    for c in cache:
        age = (today - dt.date.fromisoformat(c["checked_on"])).days
        if c["outcome"] == "matched" or age < RETRY_DAYS:
            fresh[(c["address_key"], c["latitude"][:7], c["longitude"][:8])] = c
    done = 0
    for address, lat, lon in targets:
        key = address_key(address)
        k = (key, str(lat)[:7], str(lon)[:8])
        if not key or k in fresh:
            continue
        postal, outcome = postal_for_address(address, lat, lon)
        row = {"address_key": key, "latitude": str(lat), "longitude": str(lon), "postal": postal,
               "outcome": outcome, "checked_on": today.isoformat()}
        cache = [c for c in cache if (c["address_key"], c["latitude"][:7], c["longitude"][:8]) != k] + [row]
        fresh[k] = row
        done += 1
        throttle()
    return cache, done


def add_licence_row(path):
    wb = openpyxl.load_workbook(path)
    if "Licence" not in wb.sheetnames:
        return
    ws = wb["Licence"]
    if any("OneMap" in str(c.value or "") for row in ws.iter_rows() for c in row):
        return
    ws.append(["Postal codes", ONEMAP_ATTRIBUTION + " (where noted on a listing)."])
    _safe_commit(wb, path, lambda before, after: None)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--no-lookup", action="store_true", help="do not call OneMap; use the cache only")
    p.add_argument("--workbook", default=str(WORKBOOK))
    p.add_argument("--today")
    args = p.parse_args()
    today = dt.date.fromisoformat(args.today) if args.today else dt.datetime.now(ZoneInfo(SGT_TZ)).date()
    comm = community_path(args.workbook)
    listings = read_directory(comm)[1] if comm.exists() else []
    candidates = read_csv(STATE / "candidates.csv")

    targets = [(r.get("Address"), r.get("Latitude"), r.get("Longitude"))
               for r in listings if not r.get("Postal Code") and address_key(r.get("Address"))]
    targets += [(c.get("address"), c.get("latitude"), c.get("longitude"))
                for c in candidates if not c.get("postal_code") and c.get("review_status") == "pending"
                and address_key(c.get("address"))]
    cache = load_cache()
    looked = 0
    if not args.no_lookup:
        cache, looked = lookup(targets, cache, today)
        write_csv(CACHE, CACHE_FIELDS, cache)

    lookup_table = postal_lookup(openpyxl.load_workbook(args.workbook, read_only=True))
    changes = {}
    for r in listings:
        if r.get("Postal Code"):
            continue
        postal = cached_postal(cache, r.get("Address"), r.get("Latitude"), r.get("Longitude"))
        if not postal:
            continue
        fields = {"Postal Code": postal, "Last Updated": today.isoformat()}
        fields.update({k: v for k, v in lookup_table.get(postal[:2], {}).items() if v})
        notes = str(r.get("Additional Notes") or "")
        if "OneMap" not in notes:
            fields["Additional Notes"] = (notes + " " + ONEMAP_ATTRIBUTION + ".").strip()
        changes[r["ID"]] = fields
    if changes:
        update_records(comm, changes)
        add_licence_row(comm)
    outcomes = {}
    for c in cache:
        outcomes[c["outcome"]] = outcomes.get(c["outcome"], 0) + 1
    print(f"Addresses without postal code: {len(targets)} (listings {sum(1 for r in listings if not r.get('Postal Code'))}). "
          f"OneMap lookups this run: {looked}. Cache: {outcomes}. Published listings updated: {len(changes)}.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
