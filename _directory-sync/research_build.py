"""Turn the OpenStreetMap research extract into reviewable data for Malaysia and Japan.

Reads   _directory-sync/research/osm-<CC>.json.gz   (from research_osm.py)
Writes  _directory-sync/state/candidates-<cc>.csv   every candidate, graded, with review columns
        _directory-sync/research/coverage-<cc>.csv  per state / prefecture research log
        data/community-listings-<cc>.xlsx           publishable-tier listings (with --write-data)

Every listing is "Community suggested — not verified". An OpenStreetMap halal
tag is never treated as certification, ownership or a declaration.
"""
import argparse
import csv
import datetime as dt
import gzip
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from dsync.intl import (COUNTRIES, SCOPE, TownIndex, address_for, dedupe_intl,  # noqa: E402
                        has_cjk, names_for, norm_postal_intl)
from dsync.quality import DINING_AMENITIES, DINING_SHOPS, FRESH_YEARS, SURVEY_FRESH_YEARS, is_venue  # noqa: E402
from dsync.sources import AMENITY_TO_TYPE, SHOP_TO_TYPE, lifecycle  # noqa: E402

ROOT = HERE.parent
RESEARCH = HERE / "research"
STATE = HERE / "state"

CAND_FIELDS = ["source_id", "country", "state", "city", "city_basis", "name", "local_name", "address",
               "local_address", "address_basis",
               "postal_code", "latitude", "longitude", "establishment_type", "cuisine", "halal_claim",
               "halal_scope", "certification_claim", "website", "opening_hours", "osm_last_edit",
               "check_date", "lifecycle", "place_kind", "duplicate_of", "quality_tier", "quality_notes",
               "source_url", "review_status", "review_note"]

NEW_COLUMNS = ["Local Name", "Local Address", "State / Prefecture", "City", "Halal Scope", "Muslim-Owned",
               "Opening Hours", "Source URL", "Last Verified", "Verification Notes"]

NOT_OPEN = re.compile(r"opening soon|coming soon|under construction|temporarily closed|\btutup\b|akan dibuka|閉店|休業|準備中|近日オープン", re.I)

PORK_HINT = re.compile(r"\bpork\b|\bbabi\b|bak ?kut ?teh|char ?siew|char ?siu|siu ?yuk|\bbacon\b|\bham\b|tonkatsu|\bbuta\b|豚|とんかつ|トンカツ|豬|猪肉|叉烧|叉燒|肉骨茶", re.I)

COMMUNITY_DECLARATION = "No declaration has been submitted. Contact the establishment directly."


def years_since(text, today):
    try:
        return (today - dt.date.fromisoformat(str(text)[:10])).days / 365.25
    except ValueError:
        return None


GEOCODE = HERE / "research" / "geocode-cache.json"


def load_geocode():
    return json.loads(GEOCODE.read_text(encoding="utf-8")) if GEOCODE.exists() else {}


COUNCIL = re.compile(r"^(Majlis (Bandaraya|Perbandaran|Daerah) )|\s+(City|Municipal|District|Town) Council$|^(Dewan Bandaraya) ", re.I)


def _town(a):
    t = a.get("city") or a.get("town") or a.get("village") or a.get("municipality") or ""
    return COUNCIL.sub("", t).strip()


def _compose(a, cc):
    """Readable address from a reverse-geocoder result (no state, no postcode)."""
    if cc == "JP":
        parts = [a.get("neighbourhood"), a.get("quarter"), a.get("suburb"), a.get("city_district"), _town(a)]
    else:
        street = " ".join(x for x in (a.get("house_number"), a.get("road")) if x)
        parts = [street, a.get("neighbourhood") or a.get("suburb"), _town(a)]
    out = []
    for p in parts:
        if p and p not in out:
            out.append(p)
    return ", ".join(out)


def record(el, cc, state, towns, geo=None):
    tags = el.get("tags") or {}
    lat = el.get("lat", (el.get("center") or {}).get("lat"))
    lon = el.get("lon", (el.get("center") or {}).get("lon"))
    name, local = names_for(tags, cc)
    g = (geo or {}).get(f"osm:{el['type']}/{el['id']}", {})
    en, ja = g.get("en") or {}, g.get("ja") or {}
    city, basis = tags.get("addr:city") or "", "address"
    if city and cc == "JP" and has_cjk(city) and _town(en):
        city, basis = _town(en), "map position"          # keep town names in English
    if not city and _town(en):
        city, basis = _town(en), "map position"
    if not city:
        t = towns.nearest(lat, lon)
        city, basis = (t["name"], "nearest town") if t else ("", "")
    address, address_basis = address_for(tags, cc), "OpenStreetMap address tags"
    if cc == "JP" and has_cjk(address) and _compose(en, cc):
        address, address_basis = _compose(en, cc), "map position (OpenStreetMap reverse geocoding)"
    if not address and _compose(en, cc):
        address, address_basis = _compose(en, cc), "map position (OpenStreetMap reverse geocoding)"
    local_address = ""
    if cc == "JP":
        local_address = tags.get("addr:full") if has_cjk(tags.get("addr:full") or "") else ""
        if not local_address and ja:
            local_address = "".join(x for x in (ja.get("province"), _town(ja), ja.get("city_district"), ja.get("suburb"),
                                                 ja.get("quarter"), ja.get("neighbourhood")) if x)
    postal = norm_postal_intl(tags.get("addr:postcode"), cc) or norm_postal_intl(en.get("postcode"), cc)
    claim = tags.get("diet:halal") or ("cuisine" if "halal" in (tags.get("cuisine") or "") else "")
    cert = " ".join(f"{k}={v}" for k, v in tags.items()
                    if "halal" in k and ("cert" in k or "operator" in k or "authority" in k))
    kind = (f"amenity={tags['amenity']}" if tags.get("amenity") else
            f"shop={tags['shop']}" if tags.get("shop") else "")
    return {
        "source_id": f"osm:{el['type']}/{el['id']}", "country": COUNTRIES[cc]["name"], "state": state,
        "city": city, "city_basis": basis, "name": name, "local_name": local,
        "address": address, "local_address": local_address, "address_basis": address_basis,
        "postal_code": postal,
        "latitude": lat, "longitude": lon,
        "establishment_type": AMENITY_TO_TYPE.get(tags.get("amenity")) or SHOP_TO_TYPE.get(tags.get("shop")) or "Other",
        "cuisine": ", ".join(c.replace("_", " ").title() for c in (tags.get("cuisine") or "").split(";")
                             if c and c != "halal"),
        "halal_claim": claim, "halal_scope": SCOPE.get(claim, ""), "certification_claim": cert,
        "website": tags.get("website") or tags.get("contact:website") or "",
        "opening_hours": tags.get("opening_hours") or "",
        "osm_last_edit": (el.get("timestamp") or "")[:10],
        "check_date": (tags.get("check_date") or tags.get("survey:date") or "")[:10],
        "lifecycle": lifecycle(tags), "place_kind": kind, "duplicate_of": "",
        "conflict": "diet:non-halal=" + tags["diet:non-halal"]
                    if claim and tags.get("diet:non-halal") in ("yes", "only") else "",
        "source_url": f"https://www.openstreetmap.org/{el['type']}/{el['id']}",
    }


def grade(rec, cc, today):
    """Same tiers and freshness rules as Singapore (dsync.quality), with country bounds/postcodes."""
    notes = []
    if not rec["name"].strip():
        return "rejected", "no name"
    kind = rec["place_kind"]
    if kind:
        key, _, value = kind.partition("=")
        if not ((key == "amenity" and value in DINING_AMENITIES) or (key == "shop" and value in DINING_SHOPS)):
            return "rejected", f"not a dining establishment ({kind})"
    if rec["lifecycle"]:
        return "rejected", f"reported closed ({rec['lifecycle']})"
    pork = PORK_HINT.search(f"{rec['name']} {rec.get('local_name') or ''} {rec.get('cuisine') or ''}")
    if pork:
        return "needs_check", f"name or cuisine suggests pork ({pork.group(0).strip()}); check before listing"
    if rec.get("conflict"):
        return "needs_check", f"contradictory tags ({rec['conflict']} alongside a halal tag); check before listing"
    status_word = NOT_OPEN.search(f"{rec['name']} {rec.get('local_name') or ''}")
    if status_word:
        return "needs_check", f"name suggests it is not open ({status_word.group(0).strip()})"
    b = COUNTRIES[cc]["bounds"]
    if rec["latitude"] is None or not (b[0] <= rec["latitude"] <= b[1] and b[2] <= rec["longitude"] <= b[3]):
        return "rejected", "no coordinates or outside the country"
    if rec["duplicate_of"]:
        return "needs_check", f"possible duplicate of {rec['duplicate_of']}"
    if not rec["address"] and not rec["postal_code"]:
        notes.append("no address; located by map position only")
    elif rec.get("address_basis", "").startswith("map position"):
        notes.append("address derived from map position")
    if not rec["postal_code"]:
        notes.append("no postcode")
    if rec["halal_claim"] == "limited":
        notes.append("only selected dishes reported halal")
    age, survey = years_since(rec["osm_last_edit"], today), years_since(rec["check_date"], today)
    fresh = (age is not None and age <= FRESH_YEARS) or (survey is not None and survey <= SURVEY_FRESH_YEARS)
    if not fresh:
        return "needs_check", "; ".join(notes + [f"not updated in OpenStreetMap since {rec['osm_last_edit'] or 'unknown'}"])
    return "publishable", "; ".join(notes)


def listing_row(c, today, community_id):
    venue = is_venue({"place_kind": c["place_kind"], "establishment_type": c["establishment_type"], "name": c["name"]})
    if venue:
        claim = ("OpenStreetMap contributors report that halal food is available at this food court or "
                 "food centre. This does not mean every stall is halal: check each stall's own certificate "
                 "or declaration")
    else:
        claim = {"only": "OpenStreetMap contributors report that this place serves only halal food",
                 "yes": "OpenStreetMap contributors report that halal food is available here",
                 "limited": "OpenStreetMap contributors report that some dishes here are halal; the rest of "
                            "the menu and the kitchen may not be",
                 "cuisine": "OpenStreetMap contributors list this place's cuisine as halal"}.get(
            c["halal_claim"], "Listed from OpenStreetMap community data")
    area = c["city"] if c["city"] else c["state"]
    notes = [f"Source: OpenStreetMap (© OpenStreetMap contributors, ODbL). OSM record last updated "
             f"{c['osm_last_edit'] or 'unknown'}."]
    if c["city_basis"] == "nearest town":
        notes.append(f"Town shown is the nearest town on the map ({c['city']}).")
    if c.get("address_basis", "").startswith("map position"):
        notes.append("Address derived from the map position; confirm before visiting.")
    if c["certification_claim"]:
        notes.append("OpenStreetMap also carries a certification note that has not been checked.")
    return {
        "ID": community_id(c["source_id"]), "Name": c["name"], "Country": c["country"],
        "Address": c["address"] or f"{c['city']}, {c['state']}".strip(", "),
        "Postal Code": c["postal_code"] or None,
        "Area / Prefecture Override": f"{area}, {c['state']}" if area != c["state"] else c["state"],
        "Area / Prefecture": f"{area}, {c['state']}" if area != c["state"] else c["state"],
        "Cuisine": c["cuisine"] or None, "Establishment Type": c["establishment_type"],
        "Latitude": float(c["latitude"]), "Longitude": float(c["longitude"]),
        "Listing Status": "Community suggested — not verified",
        "Basis of Listing": f"{claim}. This has not been verified, the establishment has not made a declaration, "
                            "and it is not evidence of halal certification.",
        "Certification Status": "Not verified — community-suggested",
        "Declaration Text Override": COMMUNITY_DECLARATION, "Declaration Text": COMMUNITY_DECLARATION,
        "Additional Notes": " ".join(notes), "Website URL": c["website"] or None,
        "Last Updated": today.isoformat(), "Is Demo": "FALSE",
        "Local Name": c["local_name"] or None, "Local Address": c.get("local_address") or None,
        "State / Prefecture": c["state"], "City": c["city"] or None,
        "Halal Scope": c["halal_scope"] or None, "Muslim-Owned": "Unknown",
        "Opening Hours": c["opening_hours"] or None, "Source URL": c["source_url"],
        "Last Verified": None,
        "Verification Notes": "Halal status, certification, ownership and alcohol use not verified. " + (c["quality_notes"] or ""),
    }


def write_workbook(path, rows, template):
    import openpyxl
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from dsync.workbook import ODBL_NOTICE
    src = openpyxl.load_workbook(template, read_only=True)
    header = [c.value for c in next(src.worksheets[0].iter_rows(max_row=1))] + NEW_COLUMNS
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Establishments"
    ws.append(header)
    for r in rows:
        ws.append([r.get(h) for h in header])
    last = openpyxl.utils.get_column_letter(len(header))
    t = Table(displayName="tblEstablishments", ref=f"A1:{last}{max(2, len(rows) + 1)}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)
    lic = wb.create_sheet("Licence")
    for row in ODBL_NOTICE:
        lic.append(row)
    lic.append(["Country", rows[0]["Country"] if rows else ""])
    wb.save(path)


def build(cc, today, write_data):
    from apply import community_id
    data = json.load(gzip.open(RESEARCH / f"osm-{cc}.json.gz", "rt", encoding="utf-8"))
    towns = TownIndex(data["places"])
    geo = load_geocode()
    cands, coverage = [], []
    for code, area in data["areas"].items():
        for el in area["elements"]:
            cands.append(dict(record(el, cc, area["name"], towns, geo), _area=code))
    dupes = dedupe_intl(cands)
    for c in cands:
        c["quality_tier"], c["quality_notes"] = grade(c, cc, today)
    for code, area in data["areas"].items():
        mine = [c for c in cands if c["_area"] == code]
        tiers = {t: sum(1 for c in mine if c["quality_tier"] == t)
                 for t in ("publishable", "needs_check", "rejected")}
        coverage.append({
            "country": COUNTRIES[cc]["name"], "code": code, "state_or_prefecture": area["name"],
            "osm_eating_places": area["eating_places"], "halal_tagged_records": area["halal_tagged"],
            "duplicates": sum(1 for c in mine if c["duplicate_of"]),
            "closed_or_not_dining": sum(1 for c in mine if c["quality_tier"] == "rejected"),
            "selected_items_only": sum(1 for c in mine if c["halal_claim"] == "limited"),
            "entire_menu_reported": sum(1 for c in mine if c["halal_claim"] == "only"),
            "certification_note_in_osm": sum(1 for c in mine if c["certification_claim"]),
            "needs_check": tiers["needs_check"], "ready_for_publication": tiers["publishable"],
            "towns": len({c["city"] for c in mine if c["quality_tier"] == "publishable" and c["city"]}),
        })
    with open(STATE / f"candidates-{cc.lower()}.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=CAND_FIELDS, extrasaction="ignore")
        w.writeheader()
        for c in sorted(cands, key=lambda c: (c["state"], c["city"], c["name"])):
            w.writerow({**c, "review_status": "approved" if c["quality_tier"] == "publishable" else ""})
    with open(RESEARCH / f"coverage-{cc.lower()}.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(coverage[0]))
        w.writeheader()
        w.writerows(coverage)
    rows = [listing_row(c, today, community_id) for c in cands if c["quality_tier"] == "publishable"]
    ids = [r["ID"] for r in rows]
    assert len(ids) == len(set(ids)), "duplicate IDs"
    if write_data:
        write_workbook(ROOT / "data" / f"community-listings-{cc.lower()}.xlsx", sorted(rows, key=lambda r: (r["State / Prefecture"], r["City"] or "", r["Name"])),
                       ROOT / "data" / "community-listings.xlsx")
    print(f"{cc}: {len(cands)} candidates, {dupes} duplicates, {len(rows)} publishable, "
          f"{sum(1 for c in cands if c['quality_tier'] == 'needs_check')} needs_check, "
          f"{sum(1 for c in cands if c['quality_tier'] == 'rejected')} rejected")
    return cands, rows, coverage


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("countries", nargs="*", default=["MY", "JP"])
    ap.add_argument("--write-data", action="store_true")
    ap.add_argument("--today", default=dt.date.today().isoformat())
    a = ap.parse_args()
    for cc in a.countries:
        build(cc, dt.date.fromisoformat(a.today), a.write_data)
