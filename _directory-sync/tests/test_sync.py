"""End-to-end tests. Run: python -m unittest discover -s _directory-sync/tests -v

Each test copies the tool and the real workbook into a temporary repo and runs
sync.py / promote.py as the scheduler would, using offline fixtures.
"""
import csv
import datetime as dt
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

TOOL = Path(__file__).resolve().parents[1]
REPO = TOOL.parent
WORKBOOK = REPO / "data" / "muslim-friendly-establishments.xlsx"


def node(i, name, postcode=None, lat=1.30, lon=103.85, **tags):
    t = {"name": name, "amenity": "restaurant", "diet:halal": "yes"}
    if postcode:
        t["addr:postcode"] = postcode
    t.update(tags)
    return {"type": "node", "id": i, "lat": lat, "lon": lon, "tags": t}


BASE = [
    node(1, "Nasi Lemak House", "730123", **{"addr:street": "Woodlands Avenue 1", "addr:housenumber": "12"}),
    node(2, "Prata Corner", "560456", 1.37, 103.84),
    node(3, "Burger Barn @ Tampines", "529510", 1.354, 103.944, cuisine="burger"),
    node(4, "Burger Barn @ Jurong", "609601", 1.333, 103.742, cuisine="burger"),
    node(5, "Teh Tarik Stall", None, 1.29, 103.86, amenity="food_court"),
]


class Harness:
    def __init__(self):
        self.dir = Path(tempfile.mkdtemp())
        shutil.copytree(TOOL, self.dir / "_directory-sync",
                        ignore=shutil.ignore_patterns("state", "reports", "__pycache__", "tests"))
        (self.dir / "_directory-sync" / "inbox" / "permitted").mkdir(parents=True, exist_ok=True)
        (self.dir / "data").mkdir()
        shutil.copy(WORKBOOK, self.dir / "data" / WORKBOOK.name)
        self.tool = self.dir / "_directory-sync"

    def sync(self, osm, today="2026-10-31", extra=()):
        fixture = self.dir / "fixture.json"
        fixture.write_text(json.dumps({"osm": osm}))
        return subprocess.run(
            [sys.executable, str(self.tool / "sync.py"), "--offline", str(fixture), "--today", today, *extra],
            capture_output=True, text=True)

    def candidates(self):
        with (self.tool / "state" / "candidates.csv").open() as fh:
            return {r["source_id"]: r for r in csv.DictReader(fh)}

    def log(self):
        with (self.tool / "state" / "sync-log.csv").open() as fh:
            return list(csv.DictReader(fh))

    def issues(self):
        with (self.tool / "state" / "directory-issues.csv").open() as fh:
            return list(csv.DictReader(fh))

    def cleanup(self):
        shutil.rmtree(self.dir, ignore_errors=True)


class SyncTests(unittest.TestCase):
    def setUp(self):
        self.h = Harness()

    def tearDown(self):
        self.h.cleanup()

    def test_month_end_guard(self):
        r = self.h.sync(BASE, today="2026-10-30")
        self.assertEqual(r.returncode, 0)
        self.assertIn("not the last day", r.stdout)
        self.assertFalse((self.h.tool / "state" / "candidates.csv").exists())
        r = self.h.sync(BASE, today="2026-02-28")  # non-leap February
        self.assertIn("Monthly Synchronisation Report", r.stdout)

    def test_initial_import_and_idempotent_rerun(self):
        r = self.h.sync(BASE)
        self.assertEqual(r.returncode, 0, r.stderr)
        c = self.h.candidates()
        self.assertEqual(len(c), 5)
        self.assertTrue(all(x["review_status"] == "pending" for x in c.values()))
        before = (self.h.tool / "state" / "candidates.csv").read_text()
        self.h.sync(BASE)
        self.assertEqual(len(self.h.candidates()), 5, "re-run must not duplicate records")
        self.assertEqual(before, (self.h.tool / "state" / "candidates.csv").read_text())
        self.assertEqual(self.h.log()[-1]["unchanged"], "5")
        self.assertEqual(self.h.log()[-1]["new"], "0")

    def test_branches_kept_separate(self):
        self.h.sync(BASE)
        c = self.h.candidates()
        self.assertIn("osm:node/3", c)
        self.assertIn("osm:node/4", c)
        self.assertEqual(c["osm:node/3"]["duplicate_of"], "")
        self.assertEqual(c["osm:node/4"]["duplicate_of"], "")

    def test_update_detected_and_history_kept(self):
        self.h.sync(BASE)
        changed = [dict(e, tags=dict(e["tags"])) for e in BASE]
        changed[1]["tags"]["addr:postcode"] = "560999"
        self.h.sync(changed)
        c = self.h.candidates()["osm:node/2"]
        self.assertEqual(c["postal_code"], "560999")
        self.assertIn("changed postal_code", c["change_note"])
        self.assertEqual(self.h.log()[-1]["updated"], "1")

    def test_missing_record_flagged_not_deleted(self):
        self.h.sync(BASE)
        self.h.sync(BASE[:4])  # node 5 disappears (4/5 = 80% > partial threshold)
        c = self.h.candidates()
        self.assertIn("osm:node/5", c)
        self.assertEqual(c["osm:node/5"]["missing_runs"], "1")
        self.assertEqual(self.h.log()[-1]["no_longer_found"], "1")

    def test_source_down_changes_nothing(self):
        self.h.sync(BASE)
        before = (self.h.tool / "state" / "candidates.csv").read_text()
        r = self.h.sync(None)
        self.assertEqual(r.returncode, 2)
        self.assertEqual(self.h.log()[-1]["status"], "Failed")
        self.assertEqual(before, (self.h.tool / "state" / "candidates.csv").read_text())

    def test_incomplete_source_treated_as_partial(self):
        self.h.sync(BASE)
        r = self.h.sync(BASE[:2])  # 40% of last run
        self.assertEqual(r.returncode, 0)
        self.assertEqual(self.h.log()[-1]["status"], "Partial")
        c = self.h.candidates()
        self.assertEqual(len(c), 5)
        self.assertTrue(all(x["missing_runs"] == "0" for x in c.values()))

    def test_permitted_dataset_requires_provenance_and_dedupes(self):
        inbox = self.h.tool / "inbox" / "permitted"
        (inbox / "licensed.csv").write_text(
            "establishment_name,establishment_type,certificate_number,address,postal_code,certificate_expiry\n"
            "Prata Corner,Restaurant,EERT00000001,1 Ang Mo Kio Ave 10,560456,2027-03-31\n")
        self.h.sync(BASE)
        self.assertNotIn("permitted:licensed:EERT00000001", self.h.candidates())
        report = (self.h.tool / "reports" / "2026-10.md").read_text()
        self.assertIn("NOT read", self.h.sync(BASE).stdout)
        (inbox / "PROVENANCE.md").write_text("Permission: test licence ref 123\n")
        self.h.sync(BASE)
        c = self.h.candidates()
        lic = c["permitted:licensed:EERT00000001"]
        self.assertEqual(lic["certifying_body"], "Muis (Singapore)")
        self.assertEqual(c["osm:node/2"]["duplicate_of"], lic["source_id"])
        self.assertTrue(report)

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice not installed")
    def test_directory_checks(self):
        wb = openpyxl.load_workbook(self.h.dir / "data" / WORKBOOK.name)
        ws = wb["Establishments"]
        hdr = {c.value: c.column for c in ws[1]}
        # Turn est-001 into a live, Muis-certified listing that expired.
        ws.cell(2, hdr["Is Demo"]).value = "FALSE"
        ws.cell(2, hdr["Halal Certified"]).value = "TRUE"
        ws.cell(2, hdr["Certifying Body"]).value = "Muis (Singapore)"
        ws.cell(2, hdr["Certificate Expiry"]).value = "2026-09-30"
        # est-002 live, certified, expiring within 30 days, no body recorded
        ws.cell(3, hdr["Is Demo"]).value = "FALSE"
        ws.cell(3, hdr["Halal Certified"]).value = "TRUE"
        ws.cell(3, hdr["Certificate Expiry"]).value = "2026-11-15"
        # openpyxl drops Excel's saved formula results; recalc as Excel would on save.
        staged = self.h.dir / "staged" / WORKBOOK.name
        staged.parent.mkdir()
        wb.save(staged)
        sys.path.insert(0, str(self.h.tool))
        from dsync.workbook import recalc
        recalc(staged, self.h.dir / "data")
        self.h.sync(BASE)
        kinds = {(i["directory_id"], i["issue"]) for i in self.h.issues()}
        self.assertIn(("est-001", "Certificate expired"), kinds)
        self.assertIn(("est-002", "Certificate expiring"), kinds)
        self.assertIn(("est-002", "Certification evidence missing"), kinds)

    def test_dry_run_writes_nothing(self):
        self.h.sync(BASE, extra=["--dry-run"])
        self.assertFalse((self.h.tool / "state" / "candidates.csv").exists())


@unittest.skipUnless(shutil.which("soffice"), "LibreOffice not installed")
class PromoteTests(unittest.TestCase):
    def setUp(self):
        self.h = Harness()

    def tearDown(self):
        self.h.cleanup()

    def test_promote_appends_and_preserves_existing_rows(self):
        self.h.sync(BASE)
        path = self.h.tool / "state" / "candidates.csv"
        rows = list(csv.DictReader(path.open()))
        for r in rows:
            if r["source_id"] in ("osm:node/1", "osm:node/3"):
                r["review_status"] = "approved"
        with path.open("w", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=rows[0].keys())
            w.writeheader()
            w.writerows(rows)
        before = openpyxl.load_workbook(self.h.dir / "data" / WORKBOOK.name, data_only=True)["Establishments"]
        old = [[c.value for c in row] for row in before.iter_rows(min_row=1, max_row=9)]
        r = subprocess.run([sys.executable, str(self.h.tool / "promote.py"), "--apply"],
                           capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr + r.stdout)
        after = openpyxl.load_workbook(self.h.dir / "data" / WORKBOOK.name, data_only=True)["Establishments"]
        self.assertEqual(old, [[c.value for c in row] for row in after.iter_rows(min_row=1, max_row=9)])
        hdr = {c.value: c.column - 1 for c in after[1]}
        new = [[c.value for c in row] for row in after.iter_rows(min_row=10, max_row=11)]
        self.assertEqual([n[hdr["ID"]] for n in new], ["est-009", "est-010"])
        by_name = {n[hdr["Name"]]: n for n in new}
        nasi = by_name["Nasi Lemak House"]
        self.assertEqual(nasi[hdr["Listing Status"]], "Community suggested — not verified")
        self.assertIn(nasi[hdr["Halal Certified"]], (None, ""))
        self.assertEqual(nasi[hdr["Area / Prefecture"]], "Woodlands")
        self.assertEqual(by_name["Burger Barn @ Tampines"][hdr["Area / Prefecture"]], "Tampines")
        self.assertTrue(all(n[hdr["Map Ready"]] == "Ready" for n in new))
        c = self.h.candidates()
        self.assertEqual(c["osm:node/1"]["review_status"], "promoted")
        # Next month, the promoted records are recognised as already listed.
        self.h.sync(BASE, today="2026-11-30")
        c = self.h.candidates()
        self.assertEqual(c["osm:node/1"]["match_status"], "listed")
        self.assertEqual(self.h.log()[-1]["new"], "0")


if __name__ == "__main__":
    unittest.main()
